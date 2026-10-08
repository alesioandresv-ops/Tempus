"""Tests de integración para el panel admin de reservas (Fase 7)."""

from __future__ import annotations

import datetime as dt
import io
import uuid
from collections.abc import AsyncIterator
from zoneinfo import ZoneInfo

import pytest
import pytest_asyncio
from app.db.session import TENANT_GUC
from app.models.enums import BookingEventType
from app.modules.auth.scopes import Scope
from app.modules.auth.tokens import create_access_token
from app.modules.bookings.admin import MAX_LIMITE
from app.modules.notifications.models import NotificationRequest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import (
    AsyncConnection,
    AsyncEngine,
    AsyncSession,
    create_async_engine,
)

pytestmark = pytest.mark.integration

SLUG = "negocio-admin-bookings"
BUSINESS = uuid.UUID("eeeeeeee-eeee-7eee-8eee-eeeeeeeee002")
SERVICE = uuid.UUID("eeeeeeee-eeee-7eee-8eee-eeeeeeeee003")
PROFESSIONAL = uuid.UUID("eeeeeeee-eeee-7eee-8eee-eeeeeeeee004")
CUSTOMER = uuid.UUID("eeeeeeee-eeee-7eee-8eee-eeeeeeeee005")
#: Miembro del panel. El token del fixture `admin_token` lleva este id a
#: proposito: cada walk-in escribe `created_by_user_id` (FK a `business_users`),
#: y un id inventado reventaria el flush con un IntegrityError que no dice nada.
USER = uuid.UUID("eeeeeeee-eeee-7eee-8eee-eeeeeeeee001")

TIMEZONE = "America/Argentina/Buenos_Aires"


def _cuerpo_reserva(*, starts_at: dt.datetime, telefono: str) -> dict[str, str]:
    """Cuerpo canonico del POST publico, con lo que `BookingCreateRequest` pide.

    La ruta es `/public/bookings` y el `slug` va en el cuerpo--no hay
    `/businesses/{slug}/bookings`. Y el fin **no se deduce** del servicio en
    el request: el esquema exige `ends_at` y `local_date` explicitos, que es
    lo que el motor vuelve a validar contra la duracion real.
    """
    tz = ZoneInfo(TIMEZONE)
    return {
        "slug": SLUG,
        "service_id": str(SERVICE),
        "professional_id": str(PROFESSIONAL),
        "customer_first_name": "Ana",
        "customer_last_name": "Perez",
        "customer_phone_e164": telefono,
        "starts_at": starts_at.isoformat().replace("+00:00", "Z"),
        "ends_at": (starts_at + dt.timedelta(minutes=30)).isoformat().replace("+00:00", "Z"),
        "local_date": starts_at.astimezone(tz).date().isoformat(),
    }


async def _sembrar_migration(conn: AsyncConnection) -> None:
    # `businesses` es global pero `FORCE RLS` tambien sobre el dueno, y la
    # rama `DO UPDATE` del upsert de abajo cae en `businesses_write`, que
    # exige que el GUC coincida con el negocio. El `set_config` va en
    # local, asi que no se filtra del `begin()` de este fixture.
    await conn.execute(
        text(f"SELECT set_config('{TENANT_GUC}', :tenant, true)"),
        {"tenant": str(BUSINESS)},
    )
    # `status` es parte de la prueba y no del detalle: `booking_futuro`
    # se crea por la ruta **publica**, que solo resuelve negocios
    # `active` (default de la columna: `trial`, que no se publica).
    await conn.execute(
        text(
            """
            INSERT INTO businesses (id, name, slug, timezone, status)
            VALUES (:id, 'Negocio Admin Bookings', :slug, :tz, 'active')
            ON CONFLICT (slug) DO UPDATE SET status = 'active'
            """
        ),
        {"id": BUSINESS, "slug": SLUG, "tz": TIMEZONE},
    )


async def _sembrar_app(conn: AsyncConnection) -> None:
    await conn.execute(
        text(f"SELECT set_config('{TENANT_GUC}', :tenant, true)"),
        {"tenant": str(BUSINESS)},
    )

    # Horarios Lun-Vie 09:00-12:00 y 16:00-20:00
    for weekday in range(5):
        for indice, (inicio, fin) in enumerate(
            ((dt.time(9, 0), dt.time(12, 0)), (dt.time(16, 0), dt.time(20, 0)))
        ):
            await conn.execute(
                text(
                    """
                    INSERT INTO business_hours
                        (business_id, weekday, window_index, start_time, end_time, is_open)
                    VALUES (:business, :weekday, :indice, :inicio, :fin, true)
                    ON CONFLICT (business_id, weekday, window_index) DO UPDATE SET
                        start_time = EXCLUDED.start_time,
                        end_time = EXCLUDED.end_time,
                        is_open = true
                    """
                ),
                {
                    "business": BUSINESS,
                    "weekday": weekday,
                    "indice": indice,
                    "inicio": inicio,
                    "fin": fin,
                },
            )

    # Servicio
    await conn.execute(
        text(
            """
            INSERT INTO services (id, business_id, name, duration_minutes, price, currency)
            VALUES (:id, :business, 'Corte', 30, 1500.00, 'ARS')
            ON CONFLICT (id) DO UPDATE SET
                name = EXCLUDED.name, duration_minutes = EXCLUDED.duration_minutes,
                price = EXCLUDED.price, currency = EXCLUDED.currency, is_active = true
            """
        ),
        {"id": SERVICE, "business": BUSINESS},
    )

    # Profesional
    await conn.execute(
        text(
            """
            INSERT INTO professionals (id, business_id, display_name, sort_order)
            VALUES (:id, :business, 'Profe', 1)
            ON CONFLICT (id) DO UPDATE SET
                display_name = EXCLUDED.display_name, sort_order = EXCLUDED.sort_order, is_active = true
            """
        ),
        {"id": PROFESSIONAL, "business": BUSINESS},
    )

    # Asignación servicio-profesional
    await conn.execute(
        text(
            """
            INSERT INTO professional_services (professional_id, service_id, business_id, is_active)
            VALUES (:prof, :serv, :business, true)
            ON CONFLICT (professional_id, service_id) DO UPDATE SET is_active = true
            """
        ),
        {"prof": PROFESSIONAL, "serv": SERVICE, "business": BUSINESS},
    )

    # Cliente
    await conn.execute(
        text(
            """
            INSERT INTO customers (id, business_id, first_name, last_name, phone_e164)
            VALUES (:id, :business, 'Ana', 'Perez', '+54911000000')
            ON CONFLICT (id) DO UPDATE SET first_name=EXCLUDED.first_name, last_name=EXCLUDED.last_name
            """
        ),
        {"id": CUSTOMER, "business": BUSINESS},
    )

    # Miembro del panel, con id fijo: es el `user_id` del token de
    # `admin_token`, y los walk-in lo referencian por FK al escribir
    # `created_by_user_id`.
    await conn.execute(
        text(
            """
            INSERT INTO business_users (id, business_id, email, password_hash, full_name, role, status)
            VALUES (:id, :business, 'admin@tempus.test', 'x', 'Admin', 'admin', 'active')
            ON CONFLICT (id) DO NOTHING
            """
        ),
        {"id": USER, "business": BUSINESS},
    )

    await conn.execute(text(f"SELECT set_config('{TENANT_GUC}', '', true)"))


@pytest_asyncio.fixture
async def negocio_admin(
    migration_database_url: str,
    test_database_url: str,
) -> AsyncIterator[str]:
    mig_engine: AsyncEngine = create_async_engine(
        migration_database_url, poolclass=None, echo=False
    )
    app_engine: AsyncEngine = create_async_engine(test_database_url, poolclass=None, echo=False)
    try:
        async with mig_engine.begin() as conn:
            await _sembrar_migration(conn)
        # Limpieza en **setup** y no en teardown: la ruta publica commitea de
        # verdad--sesion propia, no la del test--asi que los turnos que crea
        # `booking_futuro` sobreviven al rollback y se acumulan de corrida en
        # corrida. Sin esto, el segundo test que reserva el mismo slot en la
        # misma corrida choca con la EXCLUDE de la corrida anterior y responde
        # 409 por un motivo ajeno a lo que prueba.
        #
        # Va **antes** del seed de app: si las sobras quedaran hasta despues,
        # el cliente que la ruta publica creo para el mismo telefono reventaria
        # el upsert del seed con la unicidad `(business_id, phone_e164)`.
        # En teardown no se puede: todavia hay filas del test sin deshacer, y el
        # `DELETE` se quedaria esperando sus locks hasta el timeout.
        async with mig_engine.begin() as conn:
            await conn.execute(
                text(f"SELECT set_config('{TENANT_GUC}', :tenant, true)"),
                {"tenant": str(BUSINESS)},
            )
            # `notification_requests` se va por cascade en `bookings`.
            await conn.execute(
                text("DELETE FROM bookings WHERE business_id = :b"),
                {"b": str(BUSINESS)},
            )
            await conn.execute(
                text("DELETE FROM customers WHERE business_id = :b"),
                {"b": str(BUSINESS)},
            )
        async with app_engine.begin() as conn:
            await _sembrar_app(conn)
        yield SLUG
    finally:
        await app_engine.dispose()
        await mig_engine.dispose()


@pytest_asyncio.fixture
async def admin_token(negocio_admin: str) -> str:
    """Token de panel, emitido con la API real en vez de armado a mano.

    `create_access_token` recibe los claim y arma el payload el solo; el
    `Principal` no se construye aca porque es lo que **decodifica** el token,
    no lo que lo emite. `require_principal` no toca `business_users`, asi que
    no hace falta sembrar un miembro para que autentique: si lo que se
    quisiera probar fuera el login, seria otro test.
    """
    token = create_access_token(
        user_id=USER,
        business_id=BUSINESS,
        role="ADMIN",
        scopes=[
            Scope.BOOKINGS_READ_ANY,
            Scope.BOOKINGS_WRITE_ANY,
            Scope.BOOKINGS_WALKIN,
            Scope.BUSINESS_CONFIG_READ,
            Scope.BUSINESS_CONFIG_WRITE,
            Scope.CLIENTS_READ,
        ],
    )
    return token.token


@pytest_asyncio.fixture
async def booking_futuro(
    http_client: AsyncClient, negocio_admin: str, admin_token: str
) -> uuid.UUID:
    tz = ZoneInfo(TIMEZONE)
    # Próximo lunes 09:00 local
    base = dt.datetime.now(ZoneInfo(TIMEZONE)).date() + dt.timedelta(days=30)
    lunes = base + dt.timedelta(days=(-base.weekday()) % 7)
    local = dt.datetime(lunes.year, lunes.month, lunes.day, 9, 0, tzinfo=tz)
    starts_at = local.astimezone(dt.UTC)

    resp = await http_client.post(
        "/api/v1/public/bookings",
        json=_cuerpo_reserva(starts_at=starts_at, telefono="+54911000000"),
    )
    assert resp.status_code == 201, resp.text
    # La respuesta publica es `BookingResponse`, que trae `booking_id` y
    # no `id`: `ReservaOut` (el del panel) es el que usa `id`, y mezclar
    # los dos es como se llega a un `KeyError` que no dice nada.
    return uuid.UUID(resp.json()["booking_id"])


def _auth_header(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


class TestBookingsAdmin:
    async def test_listado_limite_max(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
    ) -> None:
        """El maximo de la pagina se acepta y se devuelve tal cual.

        El esquema del query capa `limite` con `le=200` (el `MAX_LIMITE` del
        servicio): pedir 1000 es un 422 de validacion, no un corte silencioso.
        Lo que el endpoint si acepta hasta el borde es el maximo, y lo devuelve
        en la respuesta para que el cliente sepa cual es la pagina real.
        """
        resp = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
            params={"limite": MAX_LIMITE},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["limite"] == MAX_LIMITE

    async def test_filtro_por_local_date(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
        set_tenant,
    ) -> None:
        # El panel lee por la sesion del test: sin el GUC no ve ni la reserva
        # que `booking_futuro` acaba de cerrar (RLS devuelve cero filas).
        await set_tenant(BUSINESS)

        # Cargar booking creado
        resp = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
        )
        assert resp.status_code == 200, resp.text
        items = resp.json()["items"]
        assert len(items) >= 1
        # Leer directo: el id del panel sale del recurso, no del listado.
        br = await http_client.get(
            f"/api/v1/business/reservas/{booking_futuro}",
            headers=_auth_header(admin_token),
        )
        assert br.status_code == 200, br.text
        booking = br.json()
        ld = dt.date.fromisoformat(booking["local_date"])

        # Filtro correcto
        resp = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
            params={"desde": ld.isoformat(), "hasta": ld.isoformat()},
        )
        assert resp.status_code == 200
        assert len(resp.json()["items"]) >= 1

        # Filtro incorrecto
        otro = (ld + dt.timedelta(days=2)).isoformat()
        resp = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
            params={"desde": otro, "hasta": otro},
        )
        assert resp.status_code == 200
        assert len(resp.json()["items"]) == 0

    async def test_cancelar_booking_ya_pasado_da_409(
        self,
        http_client: AsyncClient,
        session: AsyncSession,
        negocio_admin: str,
        admin_token: str,
        set_tenant,
    ) -> None:
        """Cancelar un turno que ya empezo da 409, no lo cancela.

        La ruta publica no deja reservar en el pasado (`min_lead_minutes`),
        asi que el turno pasado no se crea asi nomas: se reserva uno futuro y
        se atrasa la hora en la transaccion del test, que es la misma desde la
        que el panel va a leer el estado.
        """
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        base = dt.datetime.now(ZoneInfo(TIMEZONE)).date() + dt.timedelta(days=30)
        lunes = base + dt.timedelta(days=(-base.weekday()) % 7)
        local = dt.datetime(lunes.year, lunes.month, lunes.day, 9, 0, tzinfo=tz)

        resp = await http_client.post(
            "/api/v1/public/bookings",
            json=_cuerpo_reserva(starts_at=local.astimezone(dt.UTC), telefono="+54911999999"),
        )
        assert resp.status_code == 201, resp.text
        bid = uuid.UUID(resp.json()["booking_id"])

        ayer = dt.datetime.now(ZoneInfo(TIMEZONE)).date() - dt.timedelta(days=1)
        pasado = dt.datetime(ayer.year, ayer.month, ayer.day, 9, 0, tzinfo=tz).astimezone(dt.UTC)
        await session.execute(
            text(
                "UPDATE bookings SET starts_at = :inicio, ends_at = :fin, "
                "occupied_from = :inicio, occupied_to = :fin WHERE id = :bid"
            ),
            {
                "inicio": pasado,
                "fin": pasado + dt.timedelta(minutes=30),
                "bid": bid,
            },
        )
        await session.flush()

        # Intentar cancelar
        resp = await http_client.post(
            f"/api/v1/business/reservas/{bid}/cancelar",
            headers=_auth_header(admin_token),
            json={"motivo": "test"},
        )
        assert resp.status_code == 409, resp.text

    async def test_completar_antes_de_fin_da_409(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
        set_tenant,
    ) -> None:
        await set_tenant(BUSINESS)
        resp = await http_client.post(
            f"/api/v1/business/reservas/{booking_futuro}/estado",
            headers=_auth_header(admin_token),
            json={"estado": "completed"},
        )
        assert resp.status_code == 409, resp.text

    async def test_walkin_mismo_slot_doble_409(
        self, http_client: AsyncClient, negocio_admin: str, admin_token: str, set_tenant
    ) -> None:
        # El walk-in lee el servicio por la sesion del test: sin el GUC, la RLS
        # devuelve cero filas y responde "Servicio no encontrado" (404) en vez
        # de llegar a la EXCLUDE que este test quiere probar.
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        base = dt.datetime.now(ZoneInfo(TIMEZONE)).date() + dt.timedelta(days=30)
        lunes = base + dt.timedelta(days=(-base.weekday()) % 7)
        local = dt.datetime(lunes.year, lunes.month, lunes.day, 10, 0, tzinfo=tz)
        starts_at = local.astimezone(dt.UTC)

        r1 = await http_client.post(
            "/api/v1/business/reservas/walkin",
            headers=_auth_header(admin_token),
            json={
                "service_id": str(SERVICE),
                "professional_id": str(PROFESSIONAL),
                "starts_at": starts_at.isoformat().replace("+00:00", "Z"),
                "customer_first_name": "Marta",
                "customer_last_name": "Diaz",
                "customer_phone_e164": "+54911888888",
            },
        )
        assert r1.status_code == 201, r1.text

        r2 = await http_client.post(
            "/api/v1/business/reservas/walkin",
            headers=_auth_header(admin_token),
            json={
                "service_id": str(SERVICE),
                "professional_id": str(PROFESSIONAL),
                "starts_at": starts_at.isoformat().replace("+00:00", "Z"),
                "customer_first_name": "Marta",
                "customer_last_name": "Diaz",
                "customer_phone_e164": "+54911888888",
            },
        )
        assert r2.status_code == 409, r2.text

    async def test_walkin_no_genera_notificaciones(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        test_database_url: str,
        set_tenant,
    ) -> None:
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        base = dt.datetime.now(ZoneInfo(TIMEZONE)).date() + dt.timedelta(days=30)
        lunes = base + dt.timedelta(days=(-base.weekday()) % 7)
        local = dt.datetime(lunes.year, lunes.month, lunes.day, 11, 0, tzinfo=tz)
        starts_at = local.astimezone(dt.UTC)

        before = await self._count_notif(test_database_url, BUSINESS)
        r = await http_client.post(
            "/api/v1/business/reservas/walkin",
            headers=_auth_header(admin_token),
            json={
                "service_id": str(SERVICE),
                "professional_id": str(PROFESSIONAL),
                "starts_at": starts_at.isoformat().replace("+00:00", "Z"),
                "customer_first_name": "Sofi",
                "customer_last_name": "Lopez",
                "customer_phone_e164": "+54911777777",
            },
        )
        assert r.status_code == 201, r.text
        after = await self._count_notif(test_database_url, BUSINESS)
        assert after == before

        bid = uuid.UUID(r.json()["id"])
        ev = await http_client.get(
            f"/api/v1/business/reservas/{bid}/eventos",
            headers=_auth_header(admin_token),
        )
        assert ev.status_code == 200
        eventos = ev.json()
        assert any(e["type"] == BookingEventType.CREATED.value for e in eventos)

    async def _count_notif(self, test_database_url: str, business_id: uuid.UUID) -> int:
        """Cuenta `notification_requests` del negocio, comprometidas a la base.

        La cuenta va por una conexion aparte y con el GUC puesto: sin el, la RLS
        devuelve cero para **cualquier** negocio y el test no distinguiria un
        walk-in que no notifica de un bug que borra filas.
        """
        engine = create_async_engine(test_database_url, poolclass=None)
        try:
            async with engine.begin() as conn:
                await conn.execute(
                    text(f"SELECT set_config('{TENANT_GUC}', :t, true)"),
                    {"t": str(business_id)},
                )
                res = await conn.execute(select(func.count()).select_from(NotificationRequest))
                return int(res.scalar() or 0)
        finally:
            await engine.dispose()


def _proximo_lunes(tz: ZoneInfo) -> dt.datetime:
    """Lunes proximo a las 09:00 local, la misma convencion del resto del archivo."""
    base = dt.datetime.now(tz).date() + dt.timedelta(days=30)
    lunes = base + dt.timedelta(days=(-base.weekday()) % 7)
    return dt.datetime(lunes.year, lunes.month, lunes.day, 9, 0, tzinfo=tz)


class TestFaseC:
    """FASE C: buscador `q`, reprogramar del admin y estadisticas.

    Los tres campos de la reprogramacion van juntos (igual que el flujo
    publico): el endpoint recalcula la ocupacion con lo que se manda, y la
    EXCLUDE de la base sigue siendo la autoridad anti-solapamiento.
    """

    async def test_busqueda_q_por_nombre_y_telefono(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
        set_tenant,
    ) -> None:
        """`q` matchea parcial por nombre **y** por telefono del cliente.

        El cliente de `booking_futuro` es el del seed ("Ana Perez",
        +54911000000): buscar "ana" o el prefijo del telefono tiene que
        encontrarlo, y la misma consulta con basura devuelve cero.
        """
        await set_tenant(BUSINESS)

        por_nombre = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
            params={"q": "ana"},
        )
        assert por_nombre.status_code == 200, por_nombre.text
        assert any(i["id"] == str(booking_futuro) for i in por_nombre.json()["items"])

        por_telefono = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
            params={"q": "54911000"},
        )
        assert por_telefono.status_code == 200, por_telefono.text
        assert any(i["id"] == str(booking_futuro) for i in por_telefono.json()["items"])

        sin_azar = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
            params={"q": "zzzz-no-existe"},
        )
        assert sin_azar.status_code == 200
        assert sin_azar.json()["items"] == []

    async def test_reprogramar_mueve_y_recalcula_local_date(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
        set_tenant,
    ) -> None:
        """Reprogramar mueve el turno **en su misma fila** y recalcula la fecha local.

        `local_date` se recalcula con el timezone del negocio: mover el turno de
        un lunes a un martes 09:00 en Buenos Aires no puede quedar archivado en
        el lunes. La respuesta del endpoint es la prueba.
        """
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        lunes = _proximo_lunes(tz)
        martes = lunes + dt.timedelta(days=1)
        starts = martes.astimezone(dt.UTC)

        resp = await http_client.post(
            f"/api/v1/business/reservas/{booking_futuro}/reprogramar",
            headers=_auth_header(admin_token),
            json={
                "new_starts_at": starts.isoformat().replace("+00:00", "Z"),
                "new_ends_at": (starts + dt.timedelta(minutes=30))
                .isoformat()
                .replace("+00:00", "Z"),
                "new_duration_minutes": 30,
            },
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["starts_at"] == starts.isoformat().replace("+00:00", "Z")
        assert data["local_date"] == martes.date().isoformat()
        assert data["duration_minutes"] == 30

    async def test_reprogramar_a_slot_ocupado_da_409(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
        set_tenant,
    ) -> None:
        """La reprogramacion del panel no puede pisar otro turno.

        `booking_futuro` vive en el lunes 09:00; otro turno (walk-in) toma el
        lunes 16:00 --dentro de la grilla 09-12/16-20-- y reprogramar ahi da
        409, porque la EXCLUDE no distingue quién escribe: una reserva que se
        mueve encima de otra es un doble turno igual que dos reservas nuevas.
        """
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        lunes = _proximo_lunes(tz)
        vacio = lunes.replace(hour=16).astimezone(dt.UTC)

        rival = await http_client.post(
            "/api/v1/business/reservas/walkin",
            headers=_auth_header(admin_token),
            json={
                "service_id": str(SERVICE),
                "professional_id": str(PROFESSIONAL),
                "starts_at": vacio.isoformat().replace("+00:00", "Z"),
                "customer_first_name": "Rival",
                "customer_last_name": "Uno",
                "customer_phone_e164": "+54911666666",
            },
        )
        assert rival.status_code == 201, rival.text

        resp = await http_client.post(
            f"/api/v1/business/reservas/{booking_futuro}/reprogramar",
            headers=_auth_header(admin_token),
            json={
                "new_starts_at": vacio.isoformat().replace("+00:00", "Z"),
                "new_ends_at": (vacio + dt.timedelta(minutes=30))
                .isoformat()
                .replace("+00:00", "Z"),
                "new_duration_minutes": 30,
            },
        )
        assert resp.status_code == 409, resp.text

    async def test_estadisticas_resumen(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
        set_tenant,
    ) -> None:
        """El resumen cuenta por `local_date` y suma el `price_snapshot` real.

        En la ventana de `booking_futuro` (unica reserva confirmada del tenant
        tras el seed) el total es 1, el ingreso es el precio del servicio y la
        tasa de cancelacion es cero. Los numeros tienen que salir de la base,
        no de una suma hecha en la aplicacion.
        """
        await set_tenant(BUSINESS)

        br = await http_client.get(
            f"/api/v1/business/reservas/{booking_futuro}",
            headers=_auth_header(admin_token),
        )
        assert br.status_code == 200, br.text
        ld = br.json()["local_date"]

        resp = await http_client.get(
            "/api/v1/business/estadisticas",
            headers=_auth_header(admin_token),
            params={"desde": ld, "hasta": ld},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["desde"] == ld
        assert data["hasta"] == ld
        assert data["total_reservas"] == 1
        assert data["ingresos"] == "1500.00"
        assert data["tasa_cancelacion"] == 0.0
        profesional = next(
            (p for p in data["por_profesional"] if p["professional_id"] == str(PROFESSIONAL)),
            None,
        )
        assert profesional is not None
        assert profesional["turnos"] == 1
        assert profesional["cancelados"] == 0

    async def test_listado_sin_filtros_trae_la_reserva_publica(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        booking_futuro: uuid.UUID,
        set_tenant,
    ) -> None:
        """Regresion FASE C: el listado sin filtros ve la reserva publica del mismo negocio.

        El sintoma reportado era `/admin` con `items: []` mientras el panel del
        profesional del mismo negocio si mostraba turnos. La causa era la sesion
        del navegador (otro negocio), no el endpoint; este test fija la regresion
        explicita que esa sesion enmascaraba: una reserva creada por la ruta
        publica aparece en el listado del admin del mismo negocio (total >= 1)
        y el primer item trae el nombre real del cliente.
        """
        await set_tenant(BUSINESS)

        resp = await http_client.get(
            "/api/v1/business/reservas",
            headers=_auth_header(admin_token),
            params={"limite": MAX_LIMITE},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["total"] >= 1
        assert data["items"], "el listado sin filtros tiene que devolver la reserva publica"
        primero = data["items"][0]
        assert "Ana" in primero["cliente_nombre"], (
            "el item trae el nombre real del cliente y no un placeholder"
        )

    async def test_me_trae_el_nombre_del_negocio(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
    ) -> None:
        """`business_name` sale de la base y es el nombre real del negocio.

        El header del panel lo muestra para que un admin en el negocio
        equivocado lo note: la sesion la decide el `tid` del token, y el
        nombre hace visible esa eleccion. Tiene que salir de `businesses`,
        no de un claim del token--los claims son una cache de 15 minutos.
        """
        resp = await http_client.get(
            "/api/v1/business/me",
            headers=_auth_header(admin_token),
        )
        assert resp.status_code == 200, resp.text
        perfil = resp.json()
        assert perfil["business_id"] == str(BUSINESS)
        assert perfil["business_name"] == "Negocio Admin Bookings"


class TestFaseD:
    """FASE D (PRIORIDAD 1): historial del cliente.

    `GET /business/clientes` agrega los totales del cliente sobre sus reservas
    (count, ultima visita, gasto de confirmadas/completadas, profesional mas
    frecuente) y ordena por ultima visita real; `GET /business/clientes/{id}/
    reservas` es el historial completo para el modal del panel. Todo sale de
    agregados sobre `bookings`, sin columnas de contador en `customers`.
    """

    async def _walkin(
        self,
        http_client: AsyncClient,
        admin_token: str,
        *,
        starts_at: dt.datetime,
        telefono: str,
        nombre: str,
    ) -> uuid.UUID:
        resp = await http_client.post(
            "/api/v1/business/reservas/walkin",
            headers=_auth_header(admin_token),
            json={
                "service_id": str(SERVICE),
                "professional_id": str(PROFESSIONAL),
                "starts_at": starts_at.isoformat().replace("+00:00", "Z"),
                "customer_first_name": nombre,
                "customer_last_name": "Test",
                "customer_phone_e164": telefono,
            },
        )
        assert resp.status_code == 201, resp.text
        return uuid.UUID(resp.json()["id"])

    async def test_clientes_agrega_totales_y_ordena_por_ultima_visita(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        set_tenant,
    ) -> None:
        """3 reservas del mismo cliente: total, gasto y profesional frecuente.

        La cancelada cuenta en `total_reservas` pero no en `total_gastado`
        (solo `confirmed`/`completed`); el orden es por ultima visita real, no
        por la columna `customers.last_booking_at`, que nadie mantiene.
        """
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        lunes = _proximo_lunes(tz)

        # Ana: tres turnos el mismo lunes (09:00, 09:30, 10:00), a $1.500 c/u.
        ids = []
        for hora, minuto in ((9, 0), (9, 30), (10, 0)):
            starts = lunes.replace(hour=hora, minute=minuto).astimezone(dt.UTC)
            ids.append(
                await self._walkin(
                    http_client,
                    admin_token,
                    starts_at=starts,
                    telefono="+54911000000",
                    nombre="Ana",
                )
            )

        # Zoe: un turno mas tarde (16:00) -> queda primera en el orden.
        zoe_starts = lunes.replace(hour=16).astimezone(dt.UTC)
        await self._walkin(
            http_client, admin_token, starts_at=zoe_starts, telefono="+54911777777", nombre="Zoe"
        )

        # Cancelar el del medio de Ana: cuenta como reserva, no como gasto.
        cancel = await http_client.post(
            f"/api/v1/business/reservas/{ids[1]}/cancelar",
            headers=_auth_header(admin_token),
            json={"motivo": "test fase d"},
        )
        assert cancel.status_code == 200, cancel.text

        lista = await http_client.get(
            "/api/v1/business/clientes",
            headers=_auth_header(admin_token),
        )
        assert lista.status_code == 200, lista.text
        clientes = lista.json()

        # Zoe vino a las 16:00, Ana a las 10:00: el que vino despues arriba.
        assert [c["first_name"] for c in clientes] == ["Zoe", "Ana"]

        ana = clientes[1]
        assert ana["first_name"] == "Ana"
        assert ana["phone_e164"] == "+54911000000"
        assert ana["total_reservas"] == 3  # incluye la cancelada
        assert ana["total_gastado"] == "3000.00"  # solo confirmed/completed
        assert ana["profesional_mas_frecuente"] == "Profe"
        esperada = lunes.replace(hour=10).astimezone(dt.UTC)
        assert dt.datetime.fromisoformat(ana["ultima_reserva"].replace("Z", "+00:00")) == esperada

        zoe = clientes[0]
        assert zoe["phone_e164"] == "+54911777777"
        assert zoe["total_reservas"] == 1
        assert zoe["total_gastado"] == "1500.00"
        assert zoe["profesional_mas_frecuente"] == "Profe"

    async def test_historial_de_un_cliente_trae_sus_reservas(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        set_tenant,
    ) -> None:
        """El modal del panel: reservas del cliente, nuevas primero, con nombres.

        La fila del historial trae servicio, profesional, precio y estado, que
        es lo que pinta el drawer. Un `customer_id` de otro tenant es un 404 y
        no una lista vacia.
        """
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        lunes = _proximo_lunes(tz)
        for hora, minuto in ((9, 0), (9, 30), (10, 0)):
            starts = lunes.replace(hour=hora, minute=minuto).astimezone(dt.UTC)
            await self._walkin(
                http_client, admin_token, starts_at=starts, telefono="+54911000000", nombre="Ana"
            )

        historial = await http_client.get(
            f"/api/v1/business/clientes/{CUSTOMER}/reservas",
            headers=_auth_header(admin_token),
        )
        assert historial.status_code == 200, historial.text
        data = historial.json()
        assert data["total"] == 3
        orden = [
            dt.datetime.fromisoformat(i["starts_at"].replace("Z", "+00:00")) for i in data["items"]
        ]
        assert orden == sorted(orden, reverse=True), "las nuevas primero"
        for item in data["items"]:
            assert item["servicio_nombre"] == "Corte"
            assert item["profesional_nombre"] == "Profe"
            assert item["price_snapshot"] == "1500.00"
            assert item["status"] == "confirmed"

        ajeno = await http_client.get(
            f"/api/v1/business/clientes/{uuid.uuid4()}/reservas",
            headers=_auth_header(admin_token),
        )
        assert ajeno.status_code == 404, ajeno.text


class TestReportes:
    """FASE D (PRIORIDAD 2): reporte mensual CSV/XLSX.

    `GET /business/reportes/mensual` exporta las reservas del mes pedido y, en
    XLSX, un resumen con lo que entro (`confirmed`/`completed`) y lo que se
    cayo (`cancelled`). La ventana va por `local_date`, la misma que usa todo el
    panel: el mes es el del negocio, no el del servidor en UTC.
    """

    async def _walkin(
        self,
        http_client: AsyncClient,
        admin_token: str,
        *,
        starts_at: dt.datetime,
        telefono: str,
        nombre: str = "Ana",
    ) -> uuid.UUID:
        resp = await http_client.post(
            "/api/v1/business/reservas/walkin",
            headers=_auth_header(admin_token),
            json={
                "service_id": str(SERVICE),
                "professional_id": str(PROFESSIONAL),
                "starts_at": starts_at.isoformat().replace("+00:00", "Z"),
                "customer_first_name": nombre,
                "customer_last_name": "Test",
                "customer_phone_e164": telefono,
            },
        )
        assert resp.status_code == 201, resp.text
        return uuid.UUID(resp.json()["id"])

    async def test_csv_solo_trae_el_mes_pedido_y_404_sin_datos(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        set_tenant,
    ) -> None:
        """El CSV trae solo el mes pedido y el nombre del archivo es el del mes.

        La ventana es `local_date`: una reserva del mes siguiente no aparece en
        el archivo del mes pedido, y un mes sin reservas responde 404 en vez de
        dejar descargar un archivo vacio que el usuario veria como "todo bien".
        """
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        lunes = _proximo_lunes(tz)
        # 5 semanas despues: mismo lunes 09:00, mes garantizado distinto (35
        # dias siempre pasan de mes) y dentro del horario del negocio.
        otro_mes = lunes + dt.timedelta(days=35)
        mes_a = f"{lunes.year:04d}-{lunes.month:02d}"
        mes_b = f"{otro_mes.year:04d}-{otro_mes.month:02d}"
        assert mes_a != mes_b

        # Una reserva en cada mes, mismo cliente y mismo profesional.
        await self._walkin(
            http_client, admin_token, starts_at=lunes.astimezone(dt.UTC), telefono="+54911999999"
        )
        await self._walkin(
            http_client,
            admin_token,
            starts_at=otro_mes.astimezone(dt.UTC),
            telefono="+54911999999",
        )

        csv_a = await http_client.get(
            "/api/v1/business/reportes/mensual",
            headers=_auth_header(admin_token),
            params={"mes": mes_a, "formato": "csv"},
        )
        assert csv_a.status_code == 200, csv_a.text
        assert csv_a.headers["content-type"].startswith("text/csv")
        assert (
            csv_a.headers["content-disposition"] == f'attachment; filename="reservas_{mes_a}.csv"'
        )

        lineas = csv_a.text.splitlines()
        assert lineas[0].lstrip("\ufeff") == (
            "fecha;hora;cliente_nombre;cliente_telefono;profesional_nombre;"
            "servicio_nombre;precio;estado"
        )
        # Una sola fila de datos: la del mes pedido, y no la del otro mes.
        assert len(lineas) == 2
        celdas = lineas[1].split(";")
        assert celdas[2] == "Ana"
        assert celdas[3] == "+54911999999"
        assert celdas[4] == "Profe"
        assert celdas[5] == "Corte"
        assert celdas[6] == "1500.00"
        assert celdas[7] == "confirmed"
        # Fecha local del negocio (cae en el mes pedido) y hora local 09:00.
        assert celdas[0].startswith(mes_a)
        assert celdas[1] == "09:00"
        assert otro_mes.date().isoformat() not in csv_a.text

        # El otro mes trae su reserva y no la del primero.
        csv_b = await http_client.get(
            "/api/v1/business/reportes/mensual",
            headers=_auth_header(admin_token),
            params={"mes": mes_b, "formato": "csv"},
        )
        assert csv_b.status_code == 200, csv_b.text
        assert lunes.date().isoformat() not in csv_b.text

        # Un tercer mes sin reservas: 404, nunca un archivo vacio.
        tercero = lunes + dt.timedelta(days=70)
        vacio = await http_client.get(
            "/api/v1/business/reportes/mensual",
            headers=_auth_header(admin_token),
            params={"mes": f"{tercero.year:04d}-{tercero.month:02d}", "formato": "csv"},
        )
        assert vacio.status_code == 404, vacio.text

    async def test_xlsx_filtros_y_resumen_suman_bien(
        self,
        http_client: AsyncClient,
        negocio_admin: str,
        admin_token: str,
        set_tenant,
    ) -> None:
        """El XLSX filtra por estado y el resumen separa ingresos de canceladas.

        Tres turnos de Ana el mismo dia (09:00, 09:30, 10:00), el del medio
        cancelado: `total_reservas` cuenta los tres, `canceladas` el del medio
        y `total_ingresos` las dos confirmadas -- $1.500 c/u -- con el desglose
        por profesional y por servicio en la hoja "Resumen".
        """
        await set_tenant(BUSINESS)

        tz = ZoneInfo(TIMEZONE)
        lunes = _proximo_lunes(tz)
        mes = f"{lunes.year:04d}-{lunes.month:02d}"

        ids = []
        for hora, minuto in ((9, 0), (9, 30), (10, 0)):
            starts = lunes.replace(hour=hora, minute=minuto).astimezone(dt.UTC)
            ids.append(
                await self._walkin(
                    http_client, admin_token, starts_at=starts, telefono="+54911888888"
                )
            )

        cancel = await http_client.post(
            f"/api/v1/business/reservas/{ids[1]}/cancelar",
            headers=_auth_header(admin_token),
            json={"motivo": "test reporte mensual"},
        )
        assert cancel.status_code == 200, cancel.text

        # CSV sin filtros: encabezado + 3 filas.
        csv = await http_client.get(
            "/api/v1/business/reportes/mensual",
            headers=_auth_header(admin_token),
            params={"mes": mes, "formato": "csv"},
        )
        assert csv.status_code == 200, csv.text
        assert len(csv.text.splitlines()) == 4

        # Filtro por estado: solo la cancelada.
        canceladas = await http_client.get(
            "/api/v1/business/reportes/mensual",
            headers=_auth_header(admin_token),
            params={"mes": mes, "formato": "csv", "estado": "cancelled"},
        )
        assert canceladas.status_code == 200, canceladas.text
        solo = canceladas.text.splitlines()
        assert len(solo) == 2
        assert solo[1].split(";")[7] == "cancelled"

        # XLSX: hoja de filas + hoja de resumen con los totales correctos.
        xlsx = await http_client.get(
            "/api/v1/business/reportes/mensual",
            headers=_auth_header(admin_token),
            params={"mes": mes, "formato": "xlsx"},
        )
        assert xlsx.status_code == 200, xlsx.text
        assert (
            xlsx.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert xlsx.headers["content-disposition"] == f'attachment; filename="reservas_{mes}.xlsx"'

        libro = load_workbook(io.BytesIO(xlsx.content))
        assert libro.sheetnames == ["Reservas", "Resumen"]

        reservas = libro["Reservas"]
        assert reservas.max_row == 4  # encabezado + 3 filas
        assert reservas.max_column == 8
        assert reservas.cell(1, 3).value == "cliente_nombre"

        resumen = libro["Resumen"]
        celdas = {
            fila[0]: fila[1]
            for fila in resumen.iter_rows(values_only=True)
            if fila[0] is not None and fila[1] is not None
        }
        assert celdas["total_reservas"] == 3
        assert float(celdas["total_ingresos"]) == 3000.0
        assert celdas["canceladas"] == 1
        assert float(celdas["Profe"]) == 3000.0
        assert float(celdas["Corte"]) == 3000.0
